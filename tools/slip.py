#!/usr/bin/env python3
"""Read the job's delivery slip (<JOB>_依頼.xlsx) in one call: codes, customers, material, t, size, urgency, people.

    python tools/slip.py "//server/share/.../521030-10 Express"          # folder (finds *_依頼.xlsx) or the xlsx itself
    python tools/slip.py <folder|xlsx> --json
    python tools/slip.py <folder|xlsx> --drawing-material SPCC            # material rule check per code

One row per 作業No (the code), read from the header row 4 labels, so a moved column still works:
  customer_code (得意先, numeric, per CODE: one job lists several customers), material_text (材質), material family
  (tools/material_resolver.py: SS / SUS / AL / CU_BRASS / OTHER), surface (2B default, #1), t (板厚), size X/Y (サイズ),
  delivery date/tag (納期), creator / checker (作成担当 / チェック担当, usually empty: the DWG names say who made it).
Urgency (rules/OUTPUT.md §2): tag Express or とても急ぎます -> each code is delivered as soon as it is drawn; any other tag
(hom nay, 12h hom nay ...) -> one composite DXF.
Material rule: the slip material is used; only a different FAMILY on the drawing is a FLAG (SS vs SUS vs AL vs CU), the same
family needs no flag (slip SPHCP, drawing SPCC).
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from material_resolver import material_check, material_info  # noqa: E402,F401  (material_check re-exported)

URGENT_TAGS = ("EXPRESS", "とても急ぎます")
HEADER_ROW, FIRST_ROW = 4, 6


def norm_path(p) -> Path:
    """UNC paths work with forward slashes everywhere (a raw \\\\server\\share string fails in some Python builds)."""
    s = str(p)
    return Path(s.replace("\\", "/")) if s.startswith("\\\\") else Path(s)


def _clean(v):
    if isinstance(v, str):
        return unicodedata.normalize("NFKC", v).strip()
    return v


def find_slip(path) -> Path:
    p = norm_path(path)
    if p.is_file():
        return p
    hits = sorted(f for f in p.glob("*_依頼.xlsx") if not f.name.startswith("~$"))
    if len(hits) != 1:
        raise SystemExit(f"{p}: expected exactly one *_依頼.xlsx, found {len(hits)}")
    return hits[0]


def _columns(ws) -> dict:
    """label -> column index (1-based) from the header row; the first cell of a merged label."""
    want = {"no": "作業No", "date": "納期", "customer": "得意先", "material": "材質", "t": "板厚", "size": "サイズ",
            "creator": "作成担当", "checker": "チェック担当", "note": "備考"}
    cols = {}
    for c in ws[HEADER_ROW]:
        v = _clean(c.value)
        if not isinstance(v, str):
            continue
        for key, label in want.items():
            if label in v and key not in cols:
                cols[key] = c.column
        if v.replace(" ", "").startswith("サイズ(X2)") or v.replace(" ", "").startswith("サイズ（X2）"):
            cols["size2"] = c.column
    cols.setdefault("size2", cols.get("size", 12) + 2)
    missing = [k for k in ("no", "customer", "material", "t", "size") if k not in cols]
    if missing:
        raise SystemExit(f"slip header row {HEADER_ROW}: no column for {missing}")
    return cols


def read_slip(path) -> dict:
    f = find_slip(path)
    wb = openpyxl.load_workbook(f, data_only=True, read_only=False)
    ws = wb.active
    cols = _columns(ws)
    cell = lambda r, k, off=0: _clean(ws.cell(row=r, column=cols[k] + off).value) if k in cols else None   # noqa: E731
    date, tag = _clean(ws.cell(row=FIRST_ROW, column=cols["date"]).value), _clean(ws.cell(row=FIRST_ROW, column=cols["date"] + 1).value)
    rows = []
    for r in range(FIRST_ROW, ws.max_row + 1):
        code = cell(r, "no")
        if not code:
            continue
        mat = material_info(cell(r, "material", 1))
        mat["material_code"] = cell(r, "material")
        rows.append({
            "no": cell(r, "no", -1), "code": str(code), "customer_code": str(int(cell(r, "customer"))) if cell(r, "customer") else None,
            "customer_name": cell(r, "customer", 1), **mat, "t": cell(r, "t"),
            "size_x": cell(r, "size"), "size_y": cell(r, "size", 1), "size2_x": cell(r, "size2"), "size2_y": cell(r, "size2", 1),
            "creator": cell(r, "creator"), "checker": cell(r, "checker"), "note": cell(r, "note")})
    return {"file": str(f), "job": re.match(r"\d{6}-\d{2}", f.name).group(0) if re.match(r"\d{6}-\d{2}", f.name) else f.stem,
            "date": date, "tag": tag, "urgent": is_urgent(tag), "rows": rows}


def is_urgent(tag) -> bool:
    t = unicodedata.normalize("NFKC", str(tag or "")).upper().replace(" ", "")
    return any(u.upper().replace(" ", "") in t for u in URGENT_TAGS)


def table(slip: dict) -> str:
    lines = [f"{slip['job']}  {slip['date']} {slip['tag']}  {'URGENT: xuất từng mã' if slip['urgent'] else 'mặc định: một DXF tổng'}",
             "code | khách | vật liệu | nhóm | t | X × Y"]
    for r in slip["rows"]:
        lines.append(f"{r['code']} | {r['customer_code']} {r['customer_name'] or ''} | {r['material_text']} | "
                     f"{r['family'] or '?'}{' ' + r['surface'] if r['surface'] else ''} | {r['t']} | {r['size_x']} × {r['size_y']}")
    return "\n".join(lines)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("path")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--drawing-material", help="material printed on the drawing: report FLAG per code when the family differs")
    a = ap.parse_args(argv)
    slip = read_slip(a.path)
    if a.drawing_material:
        for r in slip["rows"]:
            r["material_check"] = material_check(r["material_text"], a.drawing_material)
    print(json.dumps(slip, ensure_ascii=False, indent=1) if a.json else table(slip))
    return 0


if __name__ == "__main__":
    for _stream in (sys.stdout, sys.stderr):
        _stream.reconfigure(encoding="utf-8")
    raise SystemExit(main())
