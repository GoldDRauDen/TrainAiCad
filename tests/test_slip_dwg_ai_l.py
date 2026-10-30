"""Slip reader, DWG roles / conversion, job layout (make/ QA/) and AI_L delivery (521030-10 session)."""
import json
import sys
from pathlib import Path

import ezdxf
import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import build_canonical as bc  # noqa: E402
import dwg_roles  # noqa: E402
import dwg_to_dxf  # noqa: E402
import job  # noqa: E402
import slip  # noqa: E402

JOB = "521030-10"
ROWS = [(1, "26-064689", 639, "有限会社橋爪鉄工所", 82, "304-2B**", 1.5, 0, 0),
        (2, "26-064687", 824, "武野工業", 12, "SPHCP(酸洗)", 1.6, 316.64, 275),
        (3, "26-064643", 795, "小田鋼機株式会社", 12, "SPHCP(酸洗)", 3.2, 152, 120)]


def make_slip(folder, tag="Express"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = JOB
    for col, text in {"C": "No.", "D": "作業No.", "E": "納期", "G": "得意先", "I": "材質", "K": "板厚", "L": "サイズ",
                      "N": "サイズ（X2）", "X": "作成担当", "Y": "チェック担当", "Z": "備考"}.items():
        ws[f"{col}4"] = text
    for i, (no, code, cust, name, mcode, mat, t, x, y) in enumerate(ROWS):
        r = 6 + i
        ws[f"C{r}"], ws[f"D{r}"], ws[f"G{r}"], ws[f"H{r}"] = no, code, cust, name
        ws[f"I{r}"], ws[f"J{r}"], ws[f"K{r}"], ws[f"L{r}"], ws[f"M{r}"] = mcode, mat, t, x, y
    ws["E6"], ws["F6"] = "10/30 (金)", tag
    path = Path(folder) / f"{JOB}_依頼.xlsx"
    wb.save(path)
    return path


def test_slip_reads_customer_per_code_material_family_and_urgency(tmp_path):
    make_slip(tmp_path)
    data = slip.read_slip(tmp_path)
    assert data["job"] == JOB and data["urgent"] is True
    r = {x["code"]: x for x in data["rows"]}
    assert [r[c]["customer_code"] for c in r] == ["639", "824", "795"]
    assert r["26-064689"]["family"] == "SUS" and r["26-064687"]["family"] == "SS"
    assert r["26-064687"]["t"] == 1.6 and (r["26-064687"]["size_x"], r["26-064687"]["size_y"]) == (316.64, 275)


@pytest.mark.parametrize("tag,urgent", [("Express", True), ("とても急ぎます", True), ("hom nay", False), ("12h hom nay", False)])
def test_urgency_tags(tag, urgent):
    assert slip.is_urgent(tag) is urgent


def test_material_rule_flags_only_a_different_family():
    assert slip.material_check("SPHCP(酸洗)", "SPCC")["flag_vi"] is None            # SS vs SS: silent
    assert slip.material_check("304-2B**", "SUS304")["flag_vi"] is None
    assert "khác" in slip.material_check("SPHCP(酸洗)", "SUS304")["flag_vi"]         # SS vs SUS
    assert "khác" in slip.material_check("A5052", "SS400")["flag_vi"]


@pytest.mark.parametrize("name,creator,checker", [
    ("521030-01-tronghieu.dwg", "tronghieu", None),
    ("521030-04_Quyen.dwg", "Quyen", None),
    ("521030-05-Toan -T.dwg", "Toan", "T"),
    ("521030-10-LUAN - hcheck.dwg", "LUAN", "h"),
    ("521030-07.dwg", "Thanh", None),
    ("521030-10-LUAN-v2.dwg", "LUAN", None),
    ("521030-12 Phat T.dwg", "Phat", "T")])
def test_dwg_names_creator_then_checker(name, creator, checker):
    info = dwg_roles.parse_name(name)
    assert info["creator"] == creator and info["checker"] == checker


def test_classify_picks_sample_and_the_matching_qa_target(tmp_path):
    for n in ("521030-10-LUAN.dwg", "521030-10-LUAN - hcheck.dwg", "521030-10-LUAN.bak", "521030-10-LUAN.dwl2"):
        (tmp_path / n).write_bytes(b"x")
    res = dwg_roles.classify(sorted(tmp_path.glob("*")), "521030-10")
    assert Path(res["sample"]).name == "521030-10-LUAN - hcheck.dwg" and Path(res["qa_target"]).name == "521030-10-LUAN.dwg"
    assert len(res["files"]) == 2          # .bak / .dwl2 are ignored


def test_dwg_to_dxf_runs_on_a_copy_and_caches(tmp_path, monkeypatch):
    dwg = tmp_path / "a.dwg"
    dwg.write_bytes(b"AC1032 fake")
    exe = tmp_path / "accoreconsole.exe"
    exe.write_bytes(b"")
    calls = []

    def runner(cmd, **kw):
        calls.append(cmd)
        assert Path(cmd[2]).name == "in.dwg" and Path(cmd[2]) != dwg and "/readonly" not in cmd
        script = Path(cmd[4]).read_bytes().decode("utf-8")   # CRLF kept
        assert script.split("\r\n")[:2] == ["FILEDIA 0", "_.DXFOUT"] and "\r\n2018\r\n" in script
        target = Path(script.split("\r\n")[2])
        doc = ezdxf.new()
        doc.modelspace().add_circle((0, 0), 5)
        doc.saveas(target)

        class R:
            stdout = b""
        return R()

    out = dwg_to_dxf.convert(dwg, tmp_path / "QA" / "a.dxf", accore=str(exe), runner=runner)
    assert ezdxf.readfile(out).modelspace()[0].dxftype() == "CIRCLE" and len(calls) == 1
    dwg_to_dxf.convert(dwg, out, accore=str(exe), runner=runner)
    assert len(calls) == 1                                     # unchanged DWG: cached
    dwg.write_bytes(b"AC1032 changed")
    dwg_to_dxf.convert(dwg, out, accore=str(exe), runner=runner)
    assert len(calls) == 2
    monkeypatch.setattr(dwg_to_dxf, "CANDIDATES", [])
    monkeypatch.delenv("ACCORECONSOLE", raising=False)
    with pytest.raises(dwg_to_dxf.DwgConvertError):
        dwg_to_dxf.find_accore()


@pytest.fixture
def jobdir(tmp_path, monkeypatch):
    monkeypatch.setattr(job, "WORK_ROOT", tmp_path / ".work")
    monkeypatch.setattr(job, "CURRENT", tmp_path / ".work" / "CURRENT")
    inbox = tmp_path / "30-10-2026" / f"{JOB} Express"
    inbox.mkdir(parents=True)
    make_slip(inbox)
    (inbox / f"{JOB}.pdf").write_bytes(b"%PDF")
    return inbox


def test_init_from_the_slip_fills_codes_customers_and_mode(jobdir):
    state = job.init(JOB, [], [str(jobdir / f"{JOB}.pdf")], None, None)
    assert list(state["codes"]) == ["26-064689", "26-064687", "26-064643"] and state["urgent"] is True
    d = job.WORK_ROOT / JOB
    census = json.loads((d / "make" / "census_26-064687.json").read_text(encoding="utf-8"))
    assert census["customer_code"] == "824" and state["codes"]["26-064643"]["slip"]["customer_code"] == "795"
    assert "URGENT" in job.summary() and "[824 SPHCP(酸洗) t1.6]" in job.summary()


def test_census_set_edits_the_census_without_hand_written_json(jobdir):
    job.init(JOB, [], [str(jobdir / f"{JOB}.pdf")], None, None)
    out = job.census(None, "26-064643", sets=["HOLE=YES:p3 4-Φ6.8"], rest="NO", source_mode="PDF_ONLY", part_mode="BEND",
                     reference="NO", multi="NO", notes=["4-Φ6.8|HOLE|ROUTED"])
    assert "rules/HOLE_TAP.md" in out["cards"]["L2_GEOMETRY"] and "rules/BEND.md" in out["cards"]["L2_GEOMETRY"]
    c = json.loads((job.WORK_ROOT / JOB / "make" / "census_26-064643.json").read_text(encoding="utf-8"))
    assert c["features"]["HOLE"]["evidence"] == ["p3 4-Φ6.8"] and c["features"]["TAP"]["state"] == "NO"
    assert c["source_manufacturing_notes"][0]["candidate_family"] == "HOLE"


def flat_job(path, codes, widen=0.0):
    parts = [{"code": c, "material": "SPCC", "t": 1.6, "envelope": [60 + i + widen, 20],
              "outline": [[0, 0], [60 + i + widen, 0], [60 + i + widen, 20], [0, 20]]} for i, c in enumerate(codes)]
    path.write_text(json.dumps(bc.build_job({"job_id": JOB, "parts": parts})), encoding="utf-8")


def test_urgent_delivery_is_one_dxf_per_code_in_ai_l_with_versions(jobdir):
    job.init(JOB, [], [str(jobdir / f"{JOB}.pdf")], None, None)
    d = job.WORK_ROOT / JOB
    codes = ["26-064689", "26-064687", "26-064643"]
    flat_job(d / "make" / f"{JOB}_job.json", codes)
    a = job.deliver_code(None, "26-064687")
    assert a == jobdir / "AI_L" / "26-064687.dxf" and len(ezdxf.readfile(a).modelspace()) > 0
    assert job.deliver_code(None, "26-064687") is None            # identical content: not repeated
    flat_job(d / "make" / f"{JOB}_job.json", codes, widen=5.0)    # redrawn: version +1
    b = job.deliver_code(None, "26-064687")
    assert b == jobdir / "AI_L" / "26-064687_v2.dxf"
    assert (jobdir / "AI_L" / "26-064687.dxf").is_file() and not list(jobdir.glob("*.dxf"))
    with pytest.raises(SystemExit):
        job.deliver_code(None, "26-999999")


def test_qa_prepare_copies_converts_and_writes_roles(jobdir, monkeypatch):
    for n in (f"{JOB}-LUAN.dwg", f"{JOB}-LUAN - hcheck.dwg"):
        (jobdir / n).write_bytes(b"AC1032 " + n.encode())
    job.init(JOB, [], [str(jobdir / f"{JOB}.pdf")], None, None)

    def fake_convert(dwg, out, **kw):
        doc = ezdxf.new()
        doc.modelspace().add_line((0, 0), (10, 0))
        doc.saveas(out)
        return Path(out)
    monkeypatch.setattr(dwg_to_dxf, "convert", fake_convert)
    roles = job.qa_prepare(None)
    qa = job.WORK_ROOT / JOB / "QA"
    assert (qa / f"{JOB}-LUAN.dwg").is_file() and (qa / f"{JOB}-LUAN - hcheck.dxf").is_file()
    assert (qa / "roles.json").is_file() and Path(roles["sample"]).name.endswith("hcheck.dwg")
    assert Path(roles["qa_target"]).name == f"{JOB}-LUAN.dwg"
