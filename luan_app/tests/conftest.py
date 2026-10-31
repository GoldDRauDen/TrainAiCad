import subprocess
from pathlib import Path

import openpyxl
import pytest

from luan_app.config import load_paths
from luan_app.store import Store


def _git(*args, cwd):
    subprocess.run(["git", *args], cwd=cwd, check=True, capture_output=True, text=True)


@pytest.fixture
def git_repo(tmp_path):
    """A repo with `main`, rules/ and tools/, pushed to a bare `origin` so `origin/main` exists."""
    origin = tmp_path / "origin.git"
    repo = tmp_path / "repo"
    _git("init", "--bare", "-b", "main", str(origin), cwd=tmp_path)
    _git("init", "-b", "main", str(repo), cwd=tmp_path)
    for k, v in {"user.email": "t@t", "user.name": "t", "core.autocrlf": "false"}.items():
        _git("config", k, v, cwd=repo)
    (repo / "rules").mkdir()
    (repo / "rules" / "X.md").write_text("x\n", encoding="utf-8")
    (repo / "tools").mkdir()
    (repo / "tools" / "noop.py").write_text("print('noop')\n", encoding="utf-8")
    _git("add", "-A", cwd=repo)
    _git("commit", "-m", "init", cwd=repo)
    _git("remote", "add", "origin", str(origin), cwd=repo)
    _git("push", "-u", "origin", "main", cwd=repo)
    return repo


@pytest.fixture
def paths(tmp_path, git_repo):
    return load_paths(data_dir=tmp_path / "data", slots_root=tmp_path / "slots", repo_root=git_repo)


@pytest.fixture
def store(paths):
    return Store(paths.db_path)


SLIP_ROWS = [(1, "26-064689", 639, "有限会社橋爪鉄工所", 82, "304-2B**", 1.5, 0, 0),
             (2, "26-064687", 824, "武野工業", 12, "SPHCP(酸洗)", 1.6, 316.64, 275),
             (3, "26-064643", 795, "小田鋼機株式会社", 12, "SPHCP(酸洗)", 3.2, 152, 120)]


def make_slip(folder, job="521030-10", tag="hom nay", rows=SLIP_ROWS):
    """Same layout tools/slip.py reads (copied from tests/test_slip_dwg_ai_l.py)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = job
    for col, text in {"C": "No.", "D": "作業No.", "E": "納期", "G": "得意先", "I": "材質", "K": "板厚", "L": "サイズ",
                      "N": "サイズ（X2）", "X": "作成担当", "Y": "チェック担当", "Z": "備考"}.items():
        ws[f"{col}4"] = text
    for i, (no, code, cust, name, mcode, mat, t, x, y) in enumerate(rows):
        r = 6 + i
        ws[f"C{r}"], ws[f"D{r}"], ws[f"G{r}"], ws[f"H{r}"] = no, code, cust, name
        ws[f"I{r}"], ws[f"J{r}"], ws[f"K{r}"], ws[f"L{r}"], ws[f"M{r}"] = mcode, mat, t, x, y
    ws["E6"], ws["F6"] = "10/30 (金)", tag
    path = Path(folder) / f"{job}_依頼.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def jobs_root(tmp_path):
    root = tmp_path / "jobs"
    root.mkdir()
    return root


@pytest.fixture
def job_dir(jobs_root):
    d = jobs_root / "521030-10"
    d.mkdir()
    (d / "521030-10.pdf").write_bytes(b"%PDF-1.4\n")
    make_slip(d)
    return d
